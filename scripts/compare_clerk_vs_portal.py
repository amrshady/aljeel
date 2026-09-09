#!/usr/bin/env python3
"""Compare an AP clerk's combined Oracle upload with portal batch exports."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re

from openpyxl import load_workbook


DEFAULT_CLERK = Path("/home/clawdbot/.openclaw/media/inbound/Main-2026-2026_Oracle-upload_3_1---3cefd2ca-6296-4ea4-bd3a-02f7ba701bf4.xlsx")
DEFAULT_PORTALS = {
    "Central 19": Path("/home/clawdbot/.openclaw/workspace/aljeel/batches/asateel-وسطي 19-2026/Central-19-2026_Oracle-upload.xlsx"),
    "Central 20": Path("/home/clawdbot/.openclaw/workspace/aljeel/batches/asateel-وسطي 20-2026/Central-20-2026_Oracle-upload.xlsx"),
}
DEFAULT_REPORT = Path("/home/clawdbot/.openclaw/workspace/aljeel/CODEX_central19-20_clerk-vs-portal_2026-09-03.md")
HEADER_ROW = 3
DATA_ROW = 4


def norm_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


HEADER_ALIASES = {
    "invoice_number": "invoice number",
    "invoice_amount": "invoice amount",
    "line_amount": "amount",
    "distribution": "distribution combination",
    "employee_no": "employee no",
    "account": "account",
    "cost_center": "cost center",
    "div": "div",
    "contribution": "contribution",
    "solution": "solution",
    "agency": "agency",
    "agency_name": "agency name",
    "location": "location",
    "line_no": "line no",
}


def header_map(ws, clerk: bool) -> dict[str, int]:
    raw = list(next(ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True)))
    found: dict[str, int] = {}
    normalized = [norm_header(x) for x in raw]
    for field, wanted in HEADER_ALIASES.items():
        matches = [i for i, h in enumerate(normalized) if h == wanted]
        if len(matches) != 1:
            raise ValueError(f"{ws.title}: expected one header for {field!r} ({wanted!r}), found {matches}")
        found[field] = matches[0]
    expected_invoice_col = 3 if clerk else 2
    if found["invoice_number"] != expected_invoice_col:
        raise ValueError(f"Unexpected invoice column in {ws.title}: {found['invoice_number']} (expected {expected_invoice_col})")
    return found


def display(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def amount(value: object) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", ""))
    except InvalidOperation as exc:
        raise ValueError(f"Invalid amount {value!r}") from exc


def comparable(field: str, value: object) -> object:
    if field in {"invoice_amount", "line_amount"}:
        return None if value in (None, "") else amount(value)
    return display(value)


def money(value: Decimal) -> str:
    return f"{value:,.2f}"


def md(value: object) -> str:
    text = display(value) if value not in (None, "") else "*(blank)*"
    return text.replace("|", "\\|").replace("\n", " ")


@dataclass
class Dataset:
    rows: dict[tuple[str, str], dict[str, object]]
    invoices: dict[str, list[dict[str, object]]]


def load_rows(path: Path, *, clerk: bool, label: str | None = None) -> Dataset:
    wb = load_workbook(path, read_only=True, data_only=True)
    if "Sheet" not in wb.sheetnames:
        raise ValueError(f"Expected sheet 'Sheet' in {path}; found {wb.sheetnames}")
    ws = wb["Sheet"]
    headers = header_map(ws, clerk)
    rows: dict[tuple[str, str], dict[str, object]] = {}
    invoices: dict[str, list[dict[str, object]]] = defaultdict(list)
    for excel_row, values in enumerate(ws.iter_rows(min_row=DATA_ROW, values_only=True), DATA_ROW):
        if clerk and display(values[0]) != label:
            continue
        invoice = display(values[headers["invoice_number"]])
        if not invoice:
            continue
        line_no = display(values[headers["line_no"]])
        if not line_no:
            raise ValueError(f"Missing Line No at {path}:{excel_row}, invoice {invoice}")
        key = (invoice, line_no)
        if key in rows:
            raise ValueError(f"Duplicate match key {key} in {path}")
        row = {field: values[index] for field, index in headers.items()}
        row["excel_row"] = excel_row
        rows[key] = row
        invoices[invoice].append(row)
    wb.close()
    return Dataset(rows, dict(invoices))


COMPARE_FIELDS = [
    "invoice_amount", "line_amount", "distribution", "employee_no", "account",
    "cost_center", "div", "contribution", "solution", "agency", "agency_name", "location",
]
FIELD_NAMES = {
    "invoice_amount": "Invoice Amount", "line_amount": "Line Amount",
    "distribution": "Distribution Combination", "employee_no": "Employee No",
    "account": "Account", "cost_center": "Cost Center", "div": "DIV",
    "contribution": "Contribution", "solution": "Solution", "agency": "Agency",
    "agency_name": "Agency Name", "location": "Location",
}
SEGMENT_NAMES = ["Company", "Location", "Account", "Cost Center", "DIV", "Solution", "Agency", "Project", "Intercompany", "Future 1"]


def compare_batch(label: str, clerk: Dataset, portal: Dataset) -> dict[str, object]:
    ck, pk = set(clerk.rows), set(portal.rows)
    missing_portal = sorted(ck - pk)
    missing_clerk = sorted(pk - ck)
    mismatches = []
    segment_mismatches = []
    counts = Counter()
    for key in sorted(ck & pk):
        c, p = clerk.rows[key], portal.rows[key]
        for field in COMPARE_FIELDS:
            if comparable(field, c[field]) != comparable(field, p[field]):
                mismatches.append((key, field, c[field], p[field]))
                counts[FIELD_NAMES[field]] += 1
        cseg, pseg = display(c["distribution"]).split("-"), display(p["distribution"]).split("-")
        for i in range(max(len(cseg), len(pseg))):
            cv = cseg[i] if i < len(cseg) else ""
            pv = pseg[i] if i < len(pseg) else ""
            if cv != pv:
                name = SEGMENT_NAMES[i] if i < len(SEGMENT_NAMES) else f"Segment {i}"
                segment_mismatches.append((key, i, name, cv, pv))
                counts[f"Distribution segment: {name}"] += 1
    structural = []
    for invoice in sorted(set(clerk.invoices) | set(portal.invoices)):
        cc, pc = len(clerk.invoices.get(invoice, [])), len(portal.invoices.get(invoice, []))
        if cc != pc:
            structural.append((invoice, cc, pc))
    return {
        "label": label, "clerk": clerk, "portal": portal,
        "missing_portal": missing_portal, "missing_clerk": missing_clerk,
        "mismatches": mismatches, "segment_mismatches": segment_mismatches,
        "structural": structural, "counts": counts,
    }


def totals(ds: Dataset) -> tuple[int, int, Decimal]:
    return len(ds.invoices), len(ds.rows), sum((amount(r["line_amount"]) for r in ds.rows.values()), Decimal("0"))


def render(results: list[dict[str, object]], clerk_path: Path, portals: dict[str, Path]) -> str:
    out = [
        "# Central 19–20 clerk vs portal Oracle upload diff",
        "",
        "**Comparison date:** 2026-09-03  ",
        f"**Clerk workbook:** `{clerk_path}`  ",
        "**Match key:** Invoice Number + explicit `Line No` (verified unique within each workbook/batch).  ",
        "**Amount handling:** exact decimal comparison and SAR totals; blank is distinct from zero for field comparisons.",
        "",
        "## Executive summary", "",
        "| Batch | Clerk invoices | Portal invoices | Clerk lines | Portal lines | Clerk line total (SAR) | Portal line total (SAR) | Missing in portal | Missing in clerk | Field mismatches | Distribution segment mismatches |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for r in results:
        ct, pt = totals(r["clerk"]), totals(r["portal"])
        out.append(f"| {r['label']} | {ct[0]} | {pt[0]} | {ct[1]} | {pt[1]} | {money(ct[2])} | {money(pt[2])} | {len(r['missing_portal'])} | {len(r['missing_clerk'])} | {len(r['mismatches'])} | {len(r['segment_mismatches'])} |")
    out += ["", "## Discrepancy categories", "", "Counts below are occurrences, not necessarily distinct lines.", ""]
    all_counts = Counter()
    for r in results:
        all_counts.update(r["counts"])
        all_counts["Lines present only in clerk"] += len(r["missing_portal"])
        all_counts["Lines present only in portal"] += len(r["missing_clerk"])
        all_counts["Invoices with differing line counts"] += len(r["structural"])
    out += ["| Category | Count |", "|---|---:|"]
    if all_counts:
        out += [f"| {md(k)} | {v} |" for k, v in sorted(all_counts.items()) if v]
    if not any(all_counts.values()):
        out.append("| No discrepancies | 0 |")

    for r in results:
        label = r["label"]
        out += ["", f"## {label}", "", f"**Portal workbook:** `{portals[label]}`", ""]
        out += ["### Structural discrepancies", ""]
        if r["structural"]:
            out += ["| Invoice | Clerk lines | Portal lines |", "|---|---:|---:|"]
            out += [f"| {md(inv)} | {cc} | {pc} |" for inv, cc, pc in r["structural"]]
        else:
            out.append("None. Invoice line counts agree.")
        out += ["", "### Lines present in clerk but missing in portal", ""]
        if r["missing_portal"]:
            out += ["| Invoice | Line |", "|---|---:|"] + [f"| {md(i)} | {md(line)} |" for i, line in r["missing_portal"]]
        else:
            out.append("None.")
        out += ["", "### Lines present in portal but missing in clerk", ""]
        if r["missing_clerk"]:
            out += ["| Invoice | Line |", "|---|---:|"] + [f"| {md(i)} | {md(line)} |" for i, line in r["missing_clerk"]]
        else:
            out.append("None.")
        out += ["", "### Per-field mismatches", ""]
        if r["mismatches"]:
            out += ["| Invoice | Line | Field | Clerk value | Portal value |", "|---|---:|---|---|---|"]
            out += [f"| {md(k[0])} | {md(k[1])} | {FIELD_NAMES[f]} | {md(cv)} | {md(pv)} |" for k, f, cv, pv in r["mismatches"]]
        else:
            out.append("None.")
        out += ["", "### Distribution Combination segment differences", ""]
        if r["segment_mismatches"]:
            out += ["| Invoice | Line | Segment index | Segment | Clerk value | Portal value |", "|---|---:|---:|---|---|---|"]
            out += [f"| {md(k[0])} | {md(k[1])} | {i} | {md(n)} | {md(cv)} | {md(pv)} |" for k, i, n, cv, pv in r["segment_mismatches"]]
        else:
            out.append("None.")
    out += ["", "## Notes", "", "Distribution segment indices are zero-based. In particular, Location is index 1 and Agency is index 6.", ""]
    return "\n".join(out)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--clerk", type=Path, default=DEFAULT_CLERK)
    parser.add_argument("--central19", type=Path, default=DEFAULT_PORTALS["Central 19"])
    parser.add_argument("--central20", type=Path, default=DEFAULT_PORTALS["Central 20"])
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()
    portals = {"Central 19": args.central19, "Central 20": args.central20}
    results = []
    for label in ("Central 19", "Central 20"):
        results.append(compare_batch(label, load_rows(args.clerk, clerk=True, label=label), load_rows(portals[label], clerk=False)))
    args.report.write_text(render(results, args.clerk, portals), encoding="utf-8")
    print("Clerk vs portal Oracle upload comparison")
    for r in results:
        ct, pt = totals(r["clerk"]), totals(r["portal"])
        print(f"{r['label']}: invoices {ct[0]}/{pt[0]}, lines {ct[1]}/{pt[1]}, line SAR {money(ct[2])}/{money(pt[2])}; only-clerk {len(r['missing_portal'])}, only-portal {len(r['missing_clerk'])}, field mismatches {len(r['mismatches'])}, segment mismatches {len(r['segment_mismatches'])}")
        for category, count in sorted(r["counts"].items()):
            print(f"  {category}: {count}")
    print(f"Report: {args.report}")


if __name__ == "__main__":
    main()
