#!/usr/bin/env python3
"""Build the human-review companion to a Jawal Oracle upload workbook."""

from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet


REFUND_HEADERS = [
    "Reference", "Ticket", "Employee No", "Employee Name", "Amount",
    "Account", "Cost Center", "DIV", "Solution", "Agency",
    "Distribution Combination", "Source Document", "Notes",
]


def _used_bounds(source: Worksheet) -> tuple[int, int]:
    """Return the last row/column containing a value (not merely formatting)."""
    max_row = max_col = 0
    for row in source.iter_rows():
        for cell in row:
            if cell.value is not None:
                max_row = max(max_row, cell.row)
                max_col = max(max_col, cell.column)
    return max_row, max_col


def _copy_sheet(source: Worksheet, target: Worksheet, *, values_only: bool = False) -> None:
    """Copy a worksheet's used value range and common presentation metadata."""
    max_row, max_col = _used_bounds(source)
    for row in source.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for source_cell in row:
            target_cell = target.cell(source_cell.row, source_cell.column, source_cell.value)
            if not values_only and source_cell.has_style:
                # Style IDs are workbook-local, so copy components rather than
                # assigning the source StyleArray across workbooks.
                target_cell.font = copy(source_cell.font)
                target_cell.fill = copy(source_cell.fill)
                target_cell.border = copy(source_cell.border)
                target_cell.number_format = source_cell.number_format
                target_cell.protection = copy(source_cell.protection)
                target_cell.alignment = copy(source_cell.alignment)

    if values_only:
        return

    for key, dimension in source.column_dimensions.items():
        target.column_dimensions[key] = copy(dimension)
    for key, dimension in source.row_dimensions.items():
        if key <= max_row:
            target.row_dimensions[key] = copy(dimension)
    for merged_range in source.merged_cells.ranges:
        if merged_range.min_row <= max_row and merged_range.min_col <= max_col:
            target.merge_cells(str(merged_range))
    target.freeze_panes = source.freeze_panes
    target.sheet_format = copy(source.sheet_format)
    target.sheet_properties = copy(source.sheet_properties)
    target.page_margins = copy(source.page_margins)
    target.page_setup = copy(source.page_setup)
    target.print_options = copy(source.print_options)
    target.auto_filter.ref = source.auto_filter.ref


def reviewer_path_for(oracle_xlsx: Path) -> Path:
    return oracle_xlsx.with_name(f"{oracle_xlsx.stem}-REVIEW.xlsx")


def build_reviewer_workbook(
    oracle_xlsx: Path,
    invoice_source_xlsx: Path,
    output_xlsx: Path | None = None,
) -> Path:
    """Create source + Oracle + refunds workbook without changing upload artifact."""
    oracle_xlsx = Path(oracle_xlsx)
    invoice_source_xlsx = Path(invoice_source_xlsx)
    output_xlsx = Path(output_xlsx) if output_xlsx else reviewer_path_for(oracle_xlsx)

    oracle_wb = openpyxl.load_workbook(oracle_xlsx, data_only=False)
    reviewer_wb = openpyxl.Workbook()
    reviewer_wb.remove(reviewer_wb.active)

    source_wb = None
    if invoice_source_xlsx.is_file():
        source_wb = openpyxl.load_workbook(invoice_source_xlsx, data_only=False)
        _copy_sheet(source_wb.worksheets[0], reviewer_wb.create_sheet("Original Invoice"))
    else:
        print(
            f"[oracle-review] WARNING: original invoice missing: {invoice_source_xlsx}; "
            "writing reviewer workbook without Original Invoice",
            flush=True,
        )

    _copy_sheet(oracle_wb.active, reviewer_wb.create_sheet("Oracle Output"))
    refunds = reviewer_wb.create_sheet("Refunds")
    if "Refunds" in oracle_wb.sheetnames:
        _copy_sheet(oracle_wb["Refunds"], refunds)
    else:
        refunds.append(REFUND_HEADERS)
        refunds.freeze_panes = "A2"
        refunds.auto_filter.ref = refunds.dimensions

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_xlsx.with_name(f".{output_xlsx.stem}.tmp.xlsx")
    reviewer_wb.save(temporary)
    temporary.replace(output_xlsx)

    if source_wb is not None:
        source_wb.close()
    oracle_wb.close()
    reviewer_wb.close()
    print(f"[oracle-review] reviewer workbook written: {output_xlsx}", flush=True)
    return output_xlsx


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("oracle_xlsx", type=Path)
    parser.add_argument("invoice_source_xlsx", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    build_reviewer_workbook(args.oracle_xlsx, args.invoice_source_xlsx, args.output)


if __name__ == "__main__":
    main()
