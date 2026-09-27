from pathlib import Path
import sys

import openpyxl
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from oracle_reviewer_workbook import build_reviewer_workbook


def _save(path: Path, title: str, values: list[list[object]]) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in values:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def test_builds_three_ordered_sheets_and_preserves_upload(tmp_path: Path) -> None:
    source = tmp_path / "invoice-source.xlsx"
    oracle = tmp_path / "Spreadsheet-J26-999-FILLED-v30.xlsx"
    _save(source, "Supplier Invoice", [["Invoice", 123], ["Formula", "=B1+1"]])
    _save(oracle, "Sheet1", [["Oracle heading"], [42]])

    source_wb = openpyxl.load_workbook(source)
    source_wb.active["A1"].font = Font(bold=True)
    source_wb.active["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    source_wb.save(source)
    source_wb.close()

    oracle_wb = openpyxl.load_workbook(oracle)
    oracle_wb.create_sheet("Refunds").append(["Ticket", "Amount"])
    oracle_wb["Refunds"].append(["T1", -50])
    oracle_wb.save(oracle)
    oracle_wb.close()

    output = build_reviewer_workbook(oracle, source)

    review_wb = openpyxl.load_workbook(output, data_only=False)
    assert review_wb.sheetnames == ["Original Invoice", "Oracle Output", "Refunds"]
    assert review_wb["Original Invoice"]["B1"].value == 123
    assert review_wb["Original Invoice"]["B2"].value == "=B1+1"
    assert review_wb["Original Invoice"]["A1"].font.bold is True
    assert review_wb["Oracle Output"]["A1"].value == "Oracle heading"
    assert review_wb["Refunds"]["B2"].value == -50
    review_wb.close()

    unchanged_wb = openpyxl.load_workbook(oracle)
    assert unchanged_wb.active.title == "Sheet1"
    assert unchanged_wb.sheetnames == ["Sheet1", "Refunds"]
    unchanged_wb.close()


def test_missing_source_warns_and_does_not_crash(tmp_path: Path, capsys) -> None:
    oracle = tmp_path / "oracle.xlsx"
    _save(oracle, "Sheet1", [["Oracle heading"]])

    output = build_reviewer_workbook(oracle, tmp_path / "missing.xlsx")

    workbook = openpyxl.load_workbook(output)
    assert workbook.sheetnames == ["Oracle Output", "Refunds"]
    assert "WARNING: original invoice missing" in capsys.readouterr().out
    workbook.close()
