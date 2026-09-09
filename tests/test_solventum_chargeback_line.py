import json
import subprocess
import sys
from collections import Counter
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from scripts.solventum_chargeback import (
    ControlledCatalogMap,
    OUTPUT_COLUMNS,
    JsonPodExtractor,
    PodExtraction,
    PodLine,
    generate_chargeback,
)


ROOT = Path(__file__).resolve().parents[1]
WAVE2 = ROOT / "batches" / "solventum" / "wave2_pods"
FIXTURE = ROOT / "tests" / "fixtures" / "solventum_wave2_pod_extractions.json"
SALES_COLUMNS = (*OUTPUT_COLUMNS, "Item Code")


class FakeExtractor:
    def __init__(self, results):
        self.results = results
        self.calls = []

    def extract(self, pdf_path, authoritative_trx):
        self.calls.append((Path(pdf_path), authoritative_trx))
        value = self.results[authoritative_trx]
        if isinstance(value, Exception):
            raise value
        return PodExtraction(authoritative_trx=authoritative_trx, lines=tuple(value),
                             document_references=("PO 9999999999", "supplier 260009137"))


def sales_row(trx, manufacturer, quantity=1, uom="Each", description=None, lot=None):
    return {
        "TRX #": trx,
        "TRX Date": "2026-06-21",
        "Order Type": "Direct_Sales",
        "Account Name": "Account",
        "Ship Address": "Riyadh",
        "Item Description": description or f"Product {manufacturer}",
        "Manufacturer": manufacturer,
        "Agency": "Solventum",
        "Lot Number": lot,
        "Quantity": quantity,
        "UOM": uom,
        "Item Code": f"ITEM-{manufacturer}",
    }


def write_sales(path, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet2"
    sheet.append(SALES_COLUMNS)
    for row in rows:
        sheet.append([row.get(column) for column in SALES_COLUMNS])
    workbook.save(path)
    workbook.close()


def read_rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(workbook["Sheet1"].iter_rows(values_only=True))
        return rows[0], [dict(zip(OUTPUT_COLUMNS, row)) for row in rows[1:]]
    finally:
        workbook.close()


def read_row_fills(path):
    workbook = load_workbook(path, data_only=True)
    try:
        return [
            [cell.fill.fgColor.rgb for cell in row[:len(OUTPUT_COLUMNS)]]
            for row in list(workbook["Sheet1"].iter_rows())[1:]
        ]
    finally:
        workbook.close()


def read_exceptions(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(workbook["Exceptions"].iter_rows(values_only=True))
        return [dict(zip(rows[0], row)) for row in rows[1:]]
    finally:
        workbook.close()


def run_line(tmp_path, sales, filenames, extractor, **kwargs):
    sales_path, output = tmp_path / "sales.xlsx", tmp_path / "output.xlsx"
    write_sales(sales_path, sales)
    pods = [tmp_path / name for name in filenames]
    count = generate_chargeback(sales_path, pods, output, match_level="line",
                                extractor=extractor, **kwargs)
    header, rows = read_rows(output)
    return count, header, rows


def test_trx_mode_filters_mixed_and_pure_non_solventum_agencies(tmp_path):
    mixed_trx = "2600019901"
    non_target_trx = "2600019943"
    solventum = sales_row(mixed_trx, "3MOC-A1")
    solventum["Agency"] = "  sOlVeNtUm  "
    ivoclar = sales_row(mixed_trx, "IVO-A1")
    ivoclar["Agency"] = "Ivoclar"
    bmx = sales_row(non_target_trx, "BMX-A1")
    bmx["Agency"] = "BMX"
    sales_path, output = tmp_path / "sales.xlsx", tmp_path / "output.xlsx"
    write_sales(sales_path, [solventum, ivoclar, bmx])

    count = generate_chargeback(
        sales_path,
        [tmp_path / f"{mixed_trx}.pdf", tmp_path / f"{non_target_trx}.pdf"],
        output,
    )

    _, rows = read_rows(output)
    assert count == 1
    assert [(row["TRX #"], row["Agency"]) for row in rows] == [
        (mixed_trx, "  sOlVeNtUm  "),
    ]


def test_line_mode_candidates_are_restricted_to_solventum_agency(tmp_path):
    trx = "2600019902"
    solventum = sales_row(trx, "A1")
    ivoclar = sales_row(trx, "A1")
    ivoclar["Agency"] = "Ivoclar"
    fake = FakeExtractor({trx: [PodLine(manufacturer_code="A1", delivered_quantity=2)]})

    count, _, rows = run_line(
        tmp_path, [solventum, ivoclar], [f"{trx}.pdf"], fake,
    )

    assert count == 1
    assert len(rows) == 1
    assert rows[0]["Agency"] == "Solventum"


def test_trx_quantity_review_highlights_every_flagged_cell_and_keeps_sales_quantities(tmp_path, capsys):
    trx = "2600010000"
    matching_trx = "2600010001"
    sales_path = tmp_path / "sales.xlsx"
    output = tmp_path / "output.xlsx"
    report = tmp_path / "review.json"
    pod = tmp_path / f"{trx}.pdf"
    write_sales(sales_path, [
        sales_row(trx, "3MOC-A1", quantity=2),
        sales_row(trx, "3MOR-A2", quantity=3),
        sales_row(matching_trx, "3MOC-B1", quantity=6),
    ])
    extractor = FakeExtractor({
        trx: [PodLine(delivered_quantity=4)],
        matching_trx: [PodLine(delivered_quantity=6)],
    })

    assert generate_chargeback(
        sales_path, [pod, tmp_path / f"{matching_trx}.pdf"], output, extractor=extractor,
        reconciliation_report=report, flag_qty_mismatch=True,
    ) == 3

    header, rows = read_rows(output)
    assert header == OUTPUT_COLUMNS
    assert [row["Quantity"] for row in rows] == [2, 3, 6]
    assert len(header) == 11
    fills = read_row_fills(output)
    assert all(color == "00FFF2CC" for row in fills[:2] for color in row)
    assert all(color == "00000000" for color in fills[2])
    assert json.loads(report.read_text(encoding="utf-8"))["quantity_review"] == [{
        "trx": trx,
        "sales_qty_total": 5.0,
        "pod_qty_total": 4.0,
        "difference": -1.0,
        "pod_files": [str(pod)],
    }]
    assert "Quantity review: 1 TRX flagged for POD/Sales quantity mismatch" in capsys.readouterr().out


def test_trx_default_path_has_no_fills_and_does_not_run_quantity_review(tmp_path):
    trx = "2600010000"
    sales_path = tmp_path / "sales.xlsx"
    output = tmp_path / "output.xlsx"
    write_sales(sales_path, [sales_row(trx, "3MOC-A1", quantity=2)])
    extractor = FakeExtractor({trx: [PodLine(delivered_quantity=99)]})

    assert generate_chargeback(
        sales_path, [tmp_path / f"{trx}.pdf"], output, extractor=extractor,
    ) == 1

    assert extractor.calls == []
    assert read_rows(output)[1][0]["Quantity"] == 2
    assert all(color == "00000000" for color in read_row_fills(output)[0])


def test_trx_quantity_review_without_extractor_skips_and_writes_empty_report(tmp_path, capsys):
    trx = "2600010000"
    sales_path = tmp_path / "sales.xlsx"
    output = tmp_path / "output.xlsx"
    report = tmp_path / "review.json"
    write_sales(sales_path, [sales_row(trx, "A1", quantity=2)])

    assert generate_chargeback(
        sales_path, [tmp_path / f"{trx}.pdf"], output,
        reconciliation_report=report, flag_qty_mismatch=True,
    ) == 1

    assert json.loads(report.read_text(encoding="utf-8"))["quantity_review"] == []
    assert all(color == "00000000" for color in read_row_fills(output)[0])
    stdout = capsys.readouterr().out
    assert "extraction is required for the quantity review; skipping review" in stdout
    assert "Quantity review: 0 TRX flagged for POD/Sales quantity mismatch" in stdout


def test_filename_trx_is_authoritative_and_nullable_fields_are_accepted(tmp_path):
    trx = "2600015875"
    fake = FakeExtractor({trx: [PodLine(manufacturer_code="A1", delivered_quantity=7)]})
    count, _, rows = run_line(tmp_path, [sales_row(trx, "3MOC-A1")],
                              [f"PO-4600162237 supplier-260009137 {trx}.pdf"], fake)
    assert count == 1
    assert rows[0]["TRX #"] == trx
    assert rows[0]["Manufacturer"] == "A1"
    assert rows[0]["Quantity"] == 7
    assert fake.calls[0][1] == trx


def test_prefix_normalization_repeated_shades_and_one_to_one_consumption(tmp_path):
    trx = "2600010001"
    sales = [
        sales_row(trx, "3MOC-1470A1", 99, description="FILTEK UNIVERSAL RESTORATIVE"),
        sales_row(trx, "3MOR-1470A2", 99, description="FILTEK UNIVERSAL RESTORATIVE"),
    ]
    # Reversed extraction order proves this is code-based/global, not row-greedy.
    fake = FakeExtractor({trx: [
        PodLine(manufacturer_code="1470A2", description="FILTEK UNIVERSAL", delivered_quantity=30),
        PodLine(manufacturer_code="3MOC-1470A1", description="FILTEK UNIVERSAL", delivered_quantity=20),
    ]})
    count, _, rows = run_line(tmp_path, sales, [f"{trx}.pdf"], fake)
    assert count == 2
    assert [(r["Manufacturer"], r["Quantity"]) for r in rows] == [("1470A2", 30), ("1470A1", 20)]


def test_exact_lot_has_priority_and_supports_4870a2_to_56921_exception(tmp_path):
    trx = "2600010002"
    lot = "0012625407-7020140871"
    sales = [
        sales_row(trx, "3MOC-4870A2", 2, "Bag-1", "Filtek One Bulk Fill", lot),
        sales_row(trx, "3MOC-56921", 2, "Bag-1", "Light cure kit", "OTHER"),
    ]
    fake = FakeExtractor({trx: [PodLine(
        catalog_code="56921", manufacturer_code="56921KIT",
        description="LIGHT CURE LUTING COMP KIT 56921 ", lot=lot,
        delivered_quantity=5, uom="kit-1",
    )]})
    count, _, rows = run_line(tmp_path, sales, [f"{trx}.pdf"], fake)
    assert count == 1
    assert rows[0]["Manufacturer"] == "56921KIT"
    assert rows[0]["Item Description"] == "Filtek One Bulk Fill"
    assert rows[0]["Lot Number"] == lot
    assert rows[0]["Quantity"] == 5
    assert rows[0]["UOM"] == "kit-1"


def test_lot_assignment_ignores_description_and_preserves_sales_description(tmp_path):
    trx = "2600010012"
    sales = [
        sales_row(trx, "A1", 1, description="Sales alpha", lot="LOT-A"),
        sales_row(trx, "B1", 1, description="Sales beta", lot="LOT-B"),
    ]
    fake = FakeExtractor({trx: [
        PodLine(manufacturer_code="B1", description="Sales alpha", lot="LOT-B", delivered_quantity=20),
        PodLine(manufacturer_code="A1", description="Sales beta", lot="LOT-A", delivered_quantity=10),
    ]})

    count, _, rows = run_line(tmp_path, sales, [f"{trx}.pdf"], fake)

    assert count == 2
    assert [(row["Lot Number"], row["Quantity"], row["Item Description"]) for row in rows] == [
        ("LOT-B", 20, "Sales beta"),
        ("LOT-A", 10, "Sales alpha"),
    ]


def test_pod_quantity_replaces_sales_and_explicit_pack_conversion_is_deterministic(tmp_path):
    trx = "2600010003"
    fake = FakeExtractor({trx: [PodLine(manufacturer_code="1954", delivered_quantity=20,
                                        uom="Box-1", pack_size=100)]})
    count, _, rows = run_line(tmp_path, [sales_row(trx, "3MOC-1954", 3000, "Box-1")],
                              [f"{trx}.pdf"], fake)
    assert count == 1
    assert rows[0]["Quantity"] == 2000
    assert rows[0]["UOM"] == "Box-1"


def test_ambiguous_extraction_can_fallback_or_error(tmp_path):
    trx = "2600010004"
    sales = [sales_row(trx, "A1", description="same product"),
             sales_row(trx, "A2", description="same product")]
    fake = FakeExtractor({trx: [PodLine(description="same product", delivered_quantity=1)]})
    with pytest.raises(ValueError, match="ambiguous"):
        run_line(tmp_path, sales, [f"{trx}.pdf"], fake, line_fallback="error")
    count, _, rows = run_line(tmp_path, sales, [f"{trx}.pdf"], fake,
                              line_fallback="trx", allow_fallback=True)
    assert count == 0
    assert rows == []


def test_multiple_filename_trx_tokens_and_repeated_page_dedup(tmp_path):
    trx1, trx2 = "2600010005", "2600010006"
    line1 = PodLine(manufacturer_code="A1", delivered_quantity=4, source_page=1)
    line1_repeat = PodLine(manufacturer_code="A1", delivered_quantity=4, source_page=2)
    fake = FakeExtractor({trx1: [line1, line1_repeat], trx2: [PodLine(manufacturer_code="B1", delivered_quantity=5)]})
    count, _, rows = run_line(
        tmp_path, [sales_row(trx1, "A1"), sales_row(trx2, "B1")],
        [f"packet {trx1}, {trx2}.pdf"], fake,
    )
    assert count == 2
    assert Counter(row["TRX #"] for row in rows) == {trx1: 1, trx2: 1}
    assert {trx for _, trx in fake.calls} == {trx1, trx2}


def test_wave2_offline_golden_trx_90_line_86_and_focal_truth(tmp_path):
    sales = WAVE2 / "JUNE SALES 2026.xlsx"
    pods = sorted(WAVE2.glob("*.pdf")) + sorted(WAVE2.glob("*.PDF"))
    trx_output, line_output = tmp_path / "trx.xlsx", tmp_path / "line.xlsx"
    assert generate_chargeback(sales, pods, trx_output) == 90
    assert generate_chargeback(sales, pods, line_output, match_level="line",
                               extractor=JsonPodExtractor(FIXTURE), allow_fallback=True) == 6
    trx_header, trx_rows = read_rows(trx_output)
    line_header, supported_rows = read_rows(line_output)
    exception_rows = read_exceptions(line_output)
    reviewed_rows = []
    for exception in exception_rows:
        # Clerk parity is supported rows plus explicitly reviewed fallback rows.
        # Unmatched Sales rows remain queued and are not deemed reviewed/claimable.
        if exception["status"] == "trx_fallback":
            reviewed_rows.extend(json.loads(exception["affected Sales rows"]))
    line_rows = supported_rows + reviewed_rows
    assert trx_header == line_header == OUTPUT_COLUMNS
    assert len(trx_rows) == 90
    assert len(supported_rows) == 6
    assert len(line_rows) == 86

    expected_counts = {
        "2600014042": 7, "2600015471": 1, "2600015669": 2, "2600015875": 6,
        "2600015947": 10, "2600015963": 7, "2600015965": 1, "2600015967": 2,
        "2600015968": 2, "2600016032": 3, "2600016033": 3, "2600016255": 1,
        "2600016282": 9, "2600016310": 4, "2600016384": 14, "2600016541": 1,
        "2600016577": 1, "2600016659": 3, "2600016726": 2, "2600016855": 1,
        "2600016942": 1, "2600017290": 2, "2600017291": 1, "2600017490": 2,
    }
    assert Counter(str(row["TRX #"]) for row in line_rows) == expected_counts
    focal = [row for row in line_rows if str(row["TRX #"]) == "2600015875"]
    assert {(str(r["Manufacturer"]), r["Quantity"], r["UOM"]) for r in focal} == {
        ("1954", 2000, "Box-1"), ("1470A3", 60, "Bag-1"),
        ("1470A2", 30, "Bag-1"), ("1470B2", 30, "Bag-1"),
        ("1470A1", 20, "Bag-1"), ("56921KIT", 5, "kit-1"),
    }
    assert {"56872", "56949", "56950", "56863"}.isdisjoint(
        {str(row["Manufacturer"]) for row in focal}
    )


def test_json_extractor_ignores_embedded_trx_and_report_records_conversion(tmp_path):
    fixture = tmp_path / "fixture.json"
    fixture.write_text(json.dumps({"extractions": {"2600010007": {
        "authoritative_trx": "WRONG-DOCUMENT-REFERENCE",
        "lines": [{"manufacturer_code": "1954", "delivered_quantity": 2, "pack_size": 100}],
    }}}), encoding="utf-8")
    report = tmp_path / "report.json"
    run_line(tmp_path, [sales_row("2600010007", "1954")], ["2600010007.pdf"],
             JsonPodExtractor(fixture), reconciliation_report=report)
    audit = json.loads(report.read_text(encoding="utf-8"))
    match = audit["transactions"]["2600010007"]["matches"][0]
    assert match["output_quantity"] == 200
    assert match["conversion"] == "delivered_quantity * pack_size(100)"


def test_controlled_normalization_alias_pack_and_raw_audit():
    normalized = ControlledCatalogMap().normalize(PodLine(
        catalog_code="CELLULOID", manufacturer_code="CELLULOID",
        description="CELLULOID 3M LIGHT CURE", delivered_quantity=20, uom="EA",
    ))
    assert normalized.raw.catalog_code == "CELLULOID"
    assert normalized.normalized.catalog_code == "1954"
    assert normalized.normalized.manufacturer_code == "1954"
    assert normalized.normalized.pack_size == 100
    assert normalized.normalized.uom == "Box-1"
    assert [rule["rule_id"] for rule in normalized.applied_rules] == [
        "alias-celluloid-soflex-to-1954", "pack-1954-box-100",
    ]


def test_controlled_identity_exception_is_lot_anchored_and_audited():
    normalized = ControlledCatalogMap().normalize(PodLine(
        catalog_code="56921", manufacturer_code="56921", delivered_quantity=5, uom="EA",
    ))
    assert normalized.normalized.manufacturer_code == "56921KIT"
    assert normalized.normalized.uom == "kit-1"
    assert normalized.match_sales_codes == ("4870A2",)
    assert normalized.match_sales_lots == ("0012625407-7020140871",)
    assert normalized.applied_rules[0]["rule_id"] == "identity-4870a2-lot-to-56921kit"


def test_trx_fallback_is_exception_only_and_requires_explicit_override(tmp_path):
    trx = "2600010008"
    sales = [sales_row(trx, "A1", description="same product"),
             sales_row(trx, "A2", description="same product")]
    report = tmp_path / "report.json"
    extractor = FakeExtractor({trx: [PodLine(description="same product", delivered_quantity=1)]})
    with pytest.raises(RuntimeError, match="--allow-fallback"):
        run_line(tmp_path, sales, [f"{trx}.pdf"], extractor, reconciliation_report=report)
    count, _, rows = run_line(
        tmp_path, sales, [f"{trx}.pdf"], extractor,
        reconciliation_report=report, allow_fallback=True,
    )
    assert count == 0
    assert rows == []
    exceptions = read_exceptions(tmp_path / "output.xlsx")
    assert len(exceptions) == 2
    assert {row["status"] for row in exceptions} == {"trx_fallback"}
    audit = json.loads(report.read_text(encoding="utf-8"))
    assert audit["transactions"][trx]["status"] == "trx_fallback"


def test_partial_reconciliation_routes_only_confident_matches_to_supported(tmp_path):
    trx = "2600010009"
    sales = [sales_row(trx, "A1"), sales_row(trx, "A2", description="same product"),
             sales_row(trx, "A3", description="same product")]
    report = tmp_path / "report.json"
    count, _, rows = run_line(
        tmp_path, sales, [f"{trx}.pdf"], FakeExtractor({trx: [
            PodLine(manufacturer_code="A1", delivered_quantity=4),
            PodLine(description="same product", delivered_quantity=1),
        ]}), line_fallback="partial", reconciliation_report=report,
    )
    assert count == 1
    assert {row["Manufacturer"] for row in rows} == {"A1"}
    exceptions = read_exceptions(tmp_path / "output.xlsx")
    assert {row["status"] for row in exceptions} == {"unmatched_sales", "unmatched_pod"}
    assert len([row for row in exceptions if row["status"] == "unmatched_sales"]) == 2
    audit = json.loads(report.read_text(encoding="utf-8"))
    assert audit["transactions"][trx]["status"] == "partial"
    assert audit["transactions"][trx]["unmatched_pod_lines"] == 1


def test_freight_line_is_excluded_without_invalidating_strong_match(tmp_path):
    trx = "2600010010"
    count, _, rows = run_line(
        tmp_path, [sales_row(trx, "A1")], [f"{trx}.pdf"], FakeExtractor({trx: [
            PodLine(manufacturer_code="A1", delivered_quantity=4),
            PodLine(description="Delivery Charges", delivered_quantity=1),
        ]}),
    )
    assert count == 1
    assert [row["Manufacturer"] for row in rows] == ["A1"]
    exceptions = read_exceptions(tmp_path / "output.xlsx")
    assert [row["status"] for row in exceptions] == ["excluded_nonchargeback_line"]


def test_cli_strict_fallback_exits_nonzero_and_override_passes_labeled(tmp_path):
    trx = "2600010011"
    sales_path = tmp_path / "sales.xlsx"
    output = tmp_path / "output.xlsx"
    extraction = tmp_path / "extraction.json"
    pod = tmp_path / f"{trx}.pdf"
    write_sales(sales_path, [
        sales_row(trx, "A1", description="same product"),
        sales_row(trx, "A2", description="same product"),
    ])
    extraction.write_text(json.dumps({"extractions": {trx: {"lines": [
        {"description": "same product", "delivered_quantity": 1},
    ]}}}), encoding="utf-8")
    command = [
        sys.executable, str(ROOT / "scripts" / "solventum_chargeback.py"),
        "--sales", str(sales_path), "--pods", str(pod), "--out", str(output),
        "--match-level", "line", "--extraction-json", str(extraction),
    ]
    strict = subprocess.run(command, cwd=ROOT, text=True, capture_output=True, check=False)
    assert strict.returncode != 0
    assert "--allow-fallback" in strict.stderr
    allowed = subprocess.run(
        [*command, "--allow-fallback"], cwd=ROOT, text=True, capture_output=True, check=False,
    )
    assert allowed.returncode == 0, allowed.stderr
    assert [row["status"] for row in read_exceptions(output)] == ["trx_fallback", "trx_fallback"]
    assert read_rows(output)[1] == []
