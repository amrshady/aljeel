import json
from pathlib import Path
import sys

import openpyxl


ROOT = Path("/home/clawdbot/.openclaw/workspace/aljeel")
sys.path.insert(0, str(ROOT / "scripts"))

import run_hybrid_v15_12 as v15
import run_v30
import split_multi_emp as splitter


HEADERS = [
    "Employee No", "*Amount", "Company", "Location", "Account", "Cost Center",
    "DIV", "Solution", "Agency", "Project", "Intercompany", "Future 1",
    "Distribution Combination", "Manpower Allocation Status", "Human Review Note",
    "Row Status", "Evidence Folder Status", "Agent Flags", "OPEX Allocation Details",
]


def _cascade_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADERS)
    ws.append([
        "", 100, "03", "40100", "60301003", "999999", "000", "00000", "00000",
        "00000", "00", "000000",
        "03-40100-60301003-999999-000-00000-00000-00000-00-000000",
        "Need to allocate", "This employee is marked 'Need to allocate' in lookups.",
        "RED", "MISSING", "FORM_NOT_FOUND_IN_EMAIL", "",
    ])
    wb.save(path)
    wb.close()


def _resolved_row(emp_no="1001663"):
    return {
        "_row_idx": 2,
        "_agent_method": "v30_authoritative_sponsoring_form",
        "emp_no": emp_no,
        "account": "60307021",
        "cost_center": "160014",
        "div": "170",
        "solution": "00000",
        "agency": "10072",
        "company": "03",
        "location": "40100",
        "project": "00000",
        "intercompany": "00",
        "future1": "000000",
        "_missing_evidence_resolved": True,
        "_sponsoring_form_folder": "/evidence/CRM-2026-44",
    }


def test_authoritative_method_writes_resolved_sponsorship_despite_form_not_found(tmp_path):
    source, output = tmp_path / "cascade.xlsx", tmp_path / "out.xlsx"
    _cascade_workbook(source)
    row = _resolved_row()
    run_v30.normalize_v30_writer_methods([row])
    v15.write_v15_12_xlsx(source, output, [row], [{h: "" for h in HEADERS}], 1)

    wb = openpyxl.load_workbook(output, data_only=True)
    ws = wb.active
    values = {ws.cell(1, col).value: ws.cell(2, col).value for col in range(1, ws.max_column + 1)}
    assert values["Employee No"] == "1001663"
    assert values["Account"] == "60307021"
    assert (values["Cost Center"], values["DIV"], values["Solution"], values["Agency"]) == (
        "160014", "170", "00000", "10072"
    )
    assert values["Distribution Combination"] == (
        "03-40100-60307021-160014-170-00000-10072-00000-00-000000"
    )
    wb.close()


def test_missing_evidence_retains_resolved_sponsorship_but_blanks_unresolved(tmp_path):
    resolved = _resolved_row()
    resolved.pop("_missing_evidence_resolved")
    unresolved = {
        "_row_idx": 3, "_agent_method": "cascade", "emp_no": "",
        "account": "60301003", "cost_center": "160014", "div": "170",
        "solution": "00000", "agency": "10072",
    }
    cascades = [
        {"Description": "NO EVIDENCE 26-901", "Invoice Ref No": ""},
        {"Description": "NO EVIDENCE 26-902", "Invoice Ref No": ""},
    ]
    assert run_v30.stamp_missing_evidence_gate([resolved, unresolved], cascades, set()) == 2
    assert resolved["_missing_evidence"] is True
    assert (resolved["emp_no"], resolved["account"], resolved["cost_center"]) == (
        "1001663", "60307021", "160014"
    )
    assert all(unresolved[key] == "" for key in (
        "emp_no", "account", "cost_center", "div", "solution", "agency"
    ))

    source, output = tmp_path / "missing-source.xlsx", tmp_path / "missing-out.xlsx"
    _cascade_workbook(source)
    run_v30.normalize_v30_writer_methods([resolved])
    v15.write_v15_12_xlsx(source, output, [resolved], [{}], 1)
    run_v30.stamp_missing_evidence_output(output, [resolved], 1)
    wb = openpyxl.load_workbook(output, data_only=True)
    ws = wb.active
    values = {ws.cell(1, col).value: ws.cell(2, col).value for col in range(1, ws.max_column + 1)}
    assert (values["Row Status"], values["Evidence Folder Status"]) == ("RED", "MISSING")
    assert (values["Employee No"], values["Account"], values["Cost Center"]) == (
        "1001663", "60307021", "160014"
    )
    wb.close()


def test_authoritative_metadata_removes_stale_need_to_allocate_note(tmp_path):
    source = tmp_path / "metadata.xlsx"
    _cascade_workbook(source)
    row = _resolved_row()
    run_v30.normalize_v30_writer_methods([row])
    run_v30.stamp_authoritative_sponsorship_metadata(
        source, [row], {"1001663": {"name": "SHEHATA"}}, 1
    )
    wb = openpyxl.load_workbook(source, data_only=True)
    ws = wb.active
    values = {ws.cell(1, col).value: ws.cell(2, col).value for col in range(1, ws.max_column + 1)}
    assert values["Manpower Allocation Status"] == "Can Be used"
    assert "Need to allocate" not in values["Human Review Note"]
    assert (values["Row Status"], values["Evidence Folder Status"]) == ("GREEN", "OK")
    wb.close()


def test_comma_separated_authoritative_sponsorship_writes_and_splits_proportionally(tmp_path, monkeypatch):
    source, written, split = tmp_path / "cascade.xlsx", tmp_path / "written.xlsx", tmp_path / "split.xlsx"
    _cascade_workbook(source)
    row = _resolved_row("1001663,1002558")
    row["_sponsorship_allocations"] = [
        {"emp_no": "1001663", "name": "SHEHATA", "amount": "3,000"},
        {"emp_no": "1002558", "name": "ALI", "amount": "1,000"},
    ]
    run_v30.normalize_v30_writer_methods([row])
    v15.write_v15_12_xlsx(source, written, [row], [{}], 1)
    run_v30.stamp_sponsorship_allocation_columns(written, [row], 1)
    monkeypatch.setattr(splitter, "load_manpower", lambda path: {})
    monkeypatch.setattr(splitter, "load_solution_names", lambda path: {})
    monkeypatch.setattr(splitter, "get_lookup", lambda path: object())
    result = splitter.split_multi_emp(str(written), str(split), "unused.xlsx")
    wb = openpyxl.load_workbook(split, data_only=True)
    ws = wb.active
    assert result["output_rows"] == 2
    assert [(ws.cell(r, 1).value, ws.cell(r, 2).value) for r in (2, 3)] == [
        ("1001663", 75), ("1002558", 25)
    ]
    wb.close()
