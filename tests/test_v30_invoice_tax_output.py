from pathlib import Path
import sys

import openpyxl


ROOT = Path("/home/clawdbot/.openclaw/workspace/aljeel")
sys.path.insert(0, str(ROOT / "scripts"))

import run_hybrid_v15_12 as writer
import split_multi_emp as splitter


def _invoice_source(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet"
    ws.cell(28, 12, "065 4861720007")
    ws.cell(28, 42, 199.57)
    ws.cell(28, 47, 1530)
    ws.cell(29, 12, "065 4861720008")
    ws.cell(29, 42, 0)
    ws.cell(29, 47, 500)
    wb.save(path)
    wb.close()


def _cascade(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.merge_cells("A2:P2")
    ws.merge_cells("Q2:AF2")
    ws.merge_cells("AG2:BO2")
    ws["A2"], ws["Q2"], ws["AG2"] = (
        "ORACLE FUSION TEMPLATE", "CODE & DESCRIPTION", "DEBUG (delete before posting)"
    )
    headers = [
        "*Invoice Header Identifier", "*Business Unit", "*Invoice Number", "*Invoice Currency",
        "*Invoice Amount", "*Invoice Date", "**Supplier[..]", "**Supplier Number",
        "*Supplier Site[..]", "Invoice Type", "Description", "*Type", "*Amount",
        "Distribution Combination[..]", "Tax Classification Code[..]", "Employee No",
        "Company",
    ]
    for col, value in enumerate(headers, 1):
        ws.cell(3, col, value)
    ws.cell(4, 5, 2000)
    ws.cell(4, 11, "BIN RAJAB/FERAS MR - JED SHW JED (4861720007)")
    ws.cell(4, 13, 1330.43)
    ws.cell(4, 14, "03-40100-60301003-160014-170-00000-10072-00000-00-000000")
    ws.cell(4, 16, "1000001,1000002")
    ws.cell(5, 11, "ZERO VAT (4861720008)")
    ws.cell(5, 13, 430)
    ws.cell(6, 11, "UNMATCHED (123)")
    ws.cell(6, 13, 99)
    ws.cell(7, 1, "Total")
    secondary = wb.create_sheet("Refunds")
    secondary["A1"] = "preserve me"
    wb.save(path)
    wb.close()


def test_v30_writer_inserts_and_populates_tax_columns_and_preserves_other_sheets(tmp_path):
    source = tmp_path / "invoice-source.xlsx"
    cascade = tmp_path / "cascade.xlsx"
    output = tmp_path / "v30.xlsx"
    _invoice_source(source)
    _cascade(cascade)

    writer.write_v15_12_xlsx(cascade, output, [], [], 3, invoice_source_path=source)

    wb = openpyxl.load_workbook(output, data_only=True)
    ws = wb["Sheet1"]
    assert [ws.cell(3, col).value for col in range(5, 7)] == [
        "*Invoice Amount", "*Invoice Date"
    ]
    assert [ws.cell(3, col).value for col in range(13, 18)] == [
        "*Amount", "VAT Amt.", "Inv. Amt. Incl. VAT",
        "Distribution Combination[..]", "Tax Classification Code[..]",
    ]
    assert (ws.cell(4, 14).value, ws.cell(4, 15).value) == (199.57, 1530)
    assert (ws.cell(5, 13).value, ws.cell(5, 14).value, ws.cell(5, 15).value) == (500, 0, 500)
    assert (ws.cell(6, 13).value, ws.cell(6, 14).value, ws.cell(6, 15).value) == (99, 0, None)
    assert (ws.cell(7, 14).value, ws.cell(7, 15).value) == (None, None)
    assert str(ws.merged_cells) == "A2:R2 S2:AH2 AI2:BQ2"
    assert wb["Refunds"]["A1"].value == "preserve me"
    wb.close()


def test_split_path_copies_new_tax_values_to_each_child(tmp_path, monkeypatch):
    source = tmp_path / "invoice-source.xlsx"
    cascade = tmp_path / "cascade.xlsx"
    written = tmp_path / "v30.xlsx"
    split = tmp_path / "v30-SPLIT.xlsx"
    _invoice_source(source)
    _cascade(cascade)
    writer.write_v15_12_xlsx(cascade, written, [], [], 3, invoice_source_path=source)
    monkeypatch.setattr(splitter, "load_manpower", lambda path: {})
    monkeypatch.setattr(splitter, "load_solution_names", lambda path: {})
    monkeypatch.setattr(splitter, "get_lookup", lambda path: object())

    splitter.split_multi_emp(str(written), str(split), "unused.xlsx")

    wb = openpyxl.load_workbook(split, data_only=True)
    ws = wb["Sheet1"]
    assert [(ws.cell(row, 14).value, ws.cell(row, 15).value) for row in (4, 5)] == [
        (199.57, 1530), (199.57, 1530)
    ]
    assert wb["Refunds"]["A1"].value == "preserve me"
    wb.close()
